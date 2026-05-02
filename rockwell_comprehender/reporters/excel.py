"""
reporters/excel.py — Reporte Excel multi-sheet del proyecto.

API pública:
    to_excel(project, output_path) -> str

Genera un archivo .xlsx con seis hojas:
    Tags     · todos los tags (todos los scopes), filtrable por scope
    Modules  · catálogo físico de módulos
    Axes     · ejes con hardware asociado y función inferida
    AOIs     · metadata de cada AOI (sin código — el código va en Markdown)
    UDTs     · UDTs con conteo de members
    Programs · programas con sus tasks asociadas

Decisiones de diseño v0.1:
- Todos los tags caben en una sola hoja con autofiltro (incluso 1500+ filas).
  El usuario filtra por scope desde Excel.
- AOIs sheet: solo metadata + conteos. El código completo de las rutinas
  internas va en el reporte Markdown — Excel no es buen soporte para código.
- Headers con estilo (fondo azul + texto blanco bold) y autofilter en cada
  hoja para usabilidad.
- Freeze de la primera fila para que el usuario navegue cómodo.
- Anchos de columnas autoajustados (cap a 60 chars para evitar columnas
  enormes con descripciones largas).
"""

from __future__ import annotations

import os
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..mapamental import (
    _build_axis_to_module,
    _infer_axis_function,
    _infer_from_tag_name,
    _is_generic_aoi,
    _looks_like_spare,
    _map_axes_to_aois,
)
from ..model import Project, Tag


# Estilo del header (consistente entre hojas)
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="305496")
_HEADER_ALIGN = Alignment(horizontal="left", vertical="center")


def to_excel(project: Project, output_path: str) -> str:
    """Genera reporte Excel multi-sheet y lo escribe a output_path.

    Args:
        project: Project ya cargado
        output_path: ruta destino del archivo .xlsx

    Returns:
        Path absoluto del archivo escrito.
    """
    abs_path = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(abs_path) or ".", exist_ok=True)

    wb = Workbook()
    # Quitar la hoja por defecto
    wb.remove(wb.active)

    _add_sheet_tags(wb, project)
    _add_sheet_modules(wb, project)
    _add_sheet_axes(wb, project)
    _add_sheet_aois(wb, project)
    _add_sheet_udts(wb, project)
    _add_sheet_programs(wb, project)

    wb.save(abs_path)
    return abs_path


# ─────────────────────────────────────────────────────────────────────────
# Hojas
# ─────────────────────────────────────────────────────────────────────────


def _add_sheet_tags(wb: Workbook, p: Project) -> None:
    headers = [
        "Name", "Scope", "DataType", "Dimension", "ExternalAccess",
        "Constant", "MotionModule", "Description",
    ]
    rows = []
    # Ordenar: controller primero, luego alfabético por scope
    for t in sorted(p.tags, key=lambda x: (x.scope != "controller", x.scope, x.name)):
        rows.append([
            t.name,
            t.scope,
            t.datatype,
            t.dimension or "",
            t.external_access or "",
            "Yes" if t.constant else "",
            t.motion_module or "",
            t.description or "",
        ])
    _write_sheet(wb, "Tags", headers, rows)


def _add_sheet_modules(wb: Workbook, p: Project) -> None:
    headers = [
        "Name", "CatalogNumber", "Vendor", "ParentModule", "ParentPortId",
        "Inhibited", "MajorFault", "ExplicitName",
    ]
    rows = []
    for m in p.modules:
        rows.append([
            m.name,
            m.catalog_number,
            m.vendor or "",
            m.parent_module or "",
            m.parent_port_id or "",
            "Yes" if m.inhibited else "",
            "Yes" if m.major_fault else "",
            "Yes" if m.has_explicit_name else "(synthetic)",
        ])
    _write_sheet(wb, "Modules", headers, rows)


def _add_sheet_axes(wb: Workbook, p: Project) -> None:
    """Hoja de ejes con todo el cruce que ya hace mapamental.

    Reusa los helpers de mapamental para no duplicar lógica.
    """
    headers = [
        "AxisName", "AxisType", "MotionModule", "ModuleCatalog", "Channel",
        "PrimaryAOI", "InferredFunction", "Description",
    ]
    rows = []

    primitive_axes = [
        t for t in p.tags
        if t.scope == "controller" and t.datatype.startswith("AXIS_")
    ]
    if not primitive_axes:
        # Hoja vacía pero presente para que el usuario sepa que se evaluó
        _write_sheet(wb, "Axes", headers, [])
        return

    axis_to_module = _build_axis_to_module(p, primitive_axes)
    axis_to_aois = _map_axes_to_aois(p, [a.name for a in primitive_axes])

    for ax in primitive_axes:
        ax_type = ax.datatype.replace("AXIS_", "")
        mod = axis_to_module.get(ax.name)
        catalog = mod.catalog_number if mod else ""
        # Canal del motion_module
        mm = (ax.motion_module or "").strip()
        channel = ""
        if ":" in mm and mm not in ("<NA>",):
            channel = mm.split(":", 1)[1]

        # Función inferida — misma lógica que mapamental
        if _looks_like_spare(ax.name):
            primary_aoi = ""
            function = "Spare / no productivo"
        elif ax.datatype == "AXIS_VIRTUAL":
            primary_aoi = ""
            function = "Eje virtual / referencia maestra"
        else:
            aois = sorted(axis_to_aois.get(ax.name, []))
            filtered_aois = [a for a in aois if not _is_generic_aoi(a)]
            primary_aoi = filtered_aois[0] if filtered_aois else (aois[0] if aois else "")
            function = (
                _infer_from_tag_name(ax.name)
                or _infer_axis_function(filtered_aois)
                or _infer_axis_function(aois)
                or ""
            )

        rows.append([
            ax.name,
            ax_type,
            mm,
            catalog,
            channel,
            primary_aoi,
            function,
            ax.description or "",
        ])

    _write_sheet(wb, "Axes", headers, rows)


def _add_sheet_aois(wb: Workbook, p: Project) -> None:
    """Solo metadata de AOIs. El código completo va en Markdown."""
    headers = [
        "Name", "Revision", "ParametersCount", "LocalTagsCount",
        "RoutinesCount", "RoutineNames", "Description",
    ]
    rows = []
    for a in sorted(p.aois, key=lambda x: x.name):
        routine_names = ", ".join(sorted(a.routines.keys())) if a.routines else ""
        rows.append([
            a.name,
            a.revision or "",
            len(a.parameters),
            len(a.local_tags),
            len(a.routines),
            routine_names,
            a.description or "",
        ])
    _write_sheet(wb, "AOIs", headers, rows)


def _add_sheet_udts(wb: Workbook, p: Project) -> None:
    headers = ["Name", "Family", "MembersCount", "Description"]
    rows = []
    for u in sorted(p.udts, key=lambda x: -len(x.members)):
        rows.append([
            u.name,
            u.family or "",
            len(u.members),
            u.description or "",
        ])
    _write_sheet(wb, "UDTs", headers, rows)


def _add_sheet_programs(wb: Workbook, p: Project) -> None:
    headers = [
        "Name", "MainRoutine", "FaultRoutine", "Disabled", "TestEdits",
        "RoutinesCount", "TagsCount", "Tasks",
    ]
    # Cruzar programas con tasks
    program_to_tasks: dict[str, list[str]] = {}
    for t in p.tasks:
        for prog_name in t.scheduled_programs:
            program_to_tasks.setdefault(prog_name, []).append(t.name)

    program_routines: dict[str, int] = {}
    for r in p.routines:
        if r.program:
            program_routines[r.program] = program_routines.get(r.program, 0) + 1
    program_tags: dict[str, int] = {}
    program_names = {pr.name for pr in p.programs}
    for t in p.tags:
        if t.scope in program_names:
            program_tags[t.scope] = program_tags.get(t.scope, 0) + 1

    rows = []
    for pr in p.programs:
        rows.append([
            pr.name,
            pr.main_routine or "",
            pr.fault_routine or "",
            "Yes" if pr.disabled else "",
            "Yes" if pr.test_edits else "",
            program_routines.get(pr.name, 0),
            program_tags.get(pr.name, 0),
            ", ".join(program_to_tasks.get(pr.name, [])),
        ])
    _write_sheet(wb, "Programs", headers, rows)


# ─────────────────────────────────────────────────────────────────────────
# Helpers de formato común
# ─────────────────────────────────────────────────────────────────────────


def _write_sheet(
    wb: Workbook,
    title: str,
    headers: list[str],
    rows: Iterable[list],
) -> Worksheet:
    """Crea una hoja con headers estilizados, filas, autofiltro y freeze."""
    ws = wb.create_sheet(title)

    # Headers
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN

    # Filas
    rows_list = list(rows)
    for ri, row in enumerate(rows_list, start=2):
        for ci, val in enumerate(row, start=1):
            ws.cell(row=ri, column=ci, value=val)

    # Auto-ancho aproximado por columna (basado en max len, capped a 60)
    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(header))
        for row in rows_list:
            if col_idx - 1 < len(row):
                v = row[col_idx - 1]
                if v is not None:
                    max_len = max(max_len, len(str(v)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 60)

    # Filtro y freeze (solo si hay datos para evitar warnings)
    if rows_list:
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{len(rows_list) + 1}"
    ws.freeze_panes = "A2"

    return ws
